import csv
import hashlib
import json
import math
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, '/tmp/slo-scale-excel-packages')
import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell as cell

OUT = Path(__file__).resolve().parent
REPORTS = OUT.parent
FINAL = REPORTS / 'final-001'
DEST = OUT / '14B_ShareGPT_SLO_0.5x_2x.xlsx'
sources = {}

def pin(path, expected=None):
    path = str(path)
    if path not in sources:
        sources[path] = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    if expected:
        assert sources[path] == expected, path
    return sources[path]

def read_csv(path):
    pin(path)
    return list(csv.DictReader(Path(path).open()))

pin(FINAL / 'REPORT.md')
rows = [r for r in read_csv(FINAL/'observations.csv') if r['measurement_host'] in ('A','C')]
boundaries = {r['measurement_host']:r for r in read_csv(FINAL/'boundaries.csv') if r['measurement_host'] in ('A','C')}
diag = [r for r in read_csv(REPORTS/'pdb-boundary-diagnostics.csv') if r['node'] in ('A','C')]
diagmap = {r['cell_id']:r for r in diag}
energy_path = REPORTS/'energy-close-final-001/energy-closeout.json'
pin(energy_path)
energy = json.loads(energy_path.read_text())
systems = ['pdblend','mixed','distserve','dynamollm','ecoserve']
names = dict(zip(systems,['PDBlend','Mixed','DistServe','DynamoLLM','EcoServe']))
colors = dict(zip(systems,['087F8C','68788F','C48220','9856A6','538B46']))
rows.sort(key=lambda r:(r['measurement_host'],float(r['rate_rps']),systems.index(r['system']),int(r['repeat'])))
assert len(rows)==62 and Counter(r['measurement_host'] for r in rows)=={'A':21,'C':41}
assert len({r['cell_id'] for r in rows})==62
assert sum(int(float(r['n_expected'])) for r in rows)==5025
assert sum(int(float(r['completed_work_requests'])) for r in rows)==5023
assert all(r['report_eligible']=='True' and r['measurement_valid']=='True' for r in rows)

evidence = []
for r in rows:
    for key in ['report_source','raw_requests','raw_power','checkpoint','receipt','summary','binding']:
        ref = json.loads(r[key])
        pin(ref['path'],ref['sha256'])
        evidence.append([float(r['slo_scale']),r['measurement_host'],names[r['system']],float(r['rate_rps']),int(r['repeat']),key,ref['path'],ref['sha256'],'已核对',r['cell_id']])
for host in ['A','C']:
    rr=[r for r in rows if r['measurement_host']==host]
    rates=sorted({float(r['rate_rps']) for r in rr})
    assert rates==[.25*(i+1) for i in range(len(rates))]
    for rate in rates:
        group=[r for r in rr if float(r['rate_rps'])==rate]
        assert len({r['trace_sha256'] for r in group})==1
        assert len([r for r in group if r['repeat']=='1'])==5
    assert math.isclose(sum(float(r['energy_j']) for r in rr),energy['nodes'][host]['primary_energy_subtotal_j'],abs_tol=1e-6)

wb=xlsxwriter.Workbook(str(DEST),{'strings_to_formulas':False,'strings_to_urls':False})
wb.set_properties({'title':'14B ShareGPT | SLO 0.5× & 2×','subject':'五系统 SLO 与能耗对比，62 次正式测量','author':'Experiment analysis','comments':'源文件、哈希、边界确认和计量范围均在工作簿内保留。'})
wb.set_calc_mode('auto')
base={'font_name':'Microsoft YaHei','font_size':10,'valign':'vcenter'}
fmt={
 'text':wb.add_format(base),
 'title':wb.add_format({**base,'font_size':20,'bold':True,'font_color':'17365D'}),
 'note':wb.add_format({**base,'font_color':'526174','text_wrap':True}),
 'section':wb.add_format({**base,'bold':True,'bg_color':'E8EFF6','font_color':'17365D'}),
 'head':wb.add_format({**base,'bold':True,'bg_color':'17365D','font_color':'FFFFFF','text_wrap':True}),
 'pct':wb.add_format({**base,'num_format':'0.00%'}),
 'num':wb.add_format({**base,'num_format':'0.000'}),
 'two':wb.add_format({**base,'num_format':'0.00'}),
 'int':wb.add_format({**base,'num_format':'0'}),
 'good':wb.add_format({**base,'bg_color':'E6F3EC','font_color':'24623C'}),
 'bad':wb.add_format({**base,'bg_color':'FCE8E6','font_color':'A12622'}),
 'confirm':wb.add_format({**base,'bg_color':'FFF1CC'}),
 'link':wb.add_format({**base,'font_color':'1767AD','underline':1}),
}
worksheets={}
def sheet(name,title,subtitle,width=22):
    ws=wb.add_worksheet(name);worksheets[name]=ws
    ws.hide_gridlines(2);ws.set_default_row(22);ws.set_column(0,25,width,fmt['text'])
    ws.merge_range(0,0,0,10,title,fmt['title']);ws.set_row(0,34)
    ws.merge_range(1,0,2,10,subtitle,fmt['note']);ws.set_row(1,22);ws.set_row(2,22)
    ws.freeze_panes(5,2);ws.set_landscape();ws.fit_to_pages(1,0);ws.repeat_rows(0,4)
    return ws

def table(ws,headers,data,formats=None):
    assert data
    ws.write_row(4,0,headers,fmt['head']);ws.set_row(4,44)
    for i,row in enumerate(data,5):
        for j,value in enumerate(row):
            ws.write(i,j,value,fmt[(formats or {}).get(j,'text')])
    ws.add_table(4,0,4+len(data),len(headers)-1,{'style':'Table Style Medium 2','columns':[{'header':h} for h in headers]})
    ws.autofit(300)
    ws.set_row(4,44)
    ws.set_column(0,0,20)

overview=sheet('说明与边界','14B ShareGPT｜SLO 0.5× 与 2×','62 次正式测量，五系统完整网格；同 trace 边界确认单独展示。源数据为 final-001 冻结报告。')
overview.set_column(0,0,24);overview.set_column(1,10,16)
head=['SLO 倍率','机器','TTFT 要求(s)','TPOT 要求(ms)','正式测量数','首测网格点/系统','最高实测≥90% rate','首次<90% rate','PDB首测达标率','PDB确认达标率','网格状态']
br=[]
for host,scale in [('A',.5),('C',2.)]:
    rr=[r for r in rows if r['measurement_host']==host and r['system']=='pdblend']
    first=next(r for r in rr if r['repeat']=='1' and float(r['slo_attainment'])<.9)
    confirm=next(r for r in rr if r['repeat']=='2')
    passed=max(float(r['rate_rps']) for r in rr if r['repeat']=='1' and float(r['slo_attainment'])>=.9)
    br.append([scale,host,float(first['slo_ttft_s']),1000*float(first['slo_tpot_s']),sum(r['measurement_host']==host for r in rows),len(rr)-1,passed,float(first['rate_rps']),float(first['slo_attainment']),float(confirm['slo_attainment']),'完整'])
table(overview,head,br,{0:'two',2:'two',3:'two',6:'two',7:'two',8:'pct',9:'pct'})
notes=[
 ('数据范围','Qwen2.5-14B-Instruct + ShareGPT；A=120.79.123.62，0.5×；C=47.106.163.29，2×。'),
 ('固定协议','arrival seed=701；内容采样种子=20260907；到达窗口100秒；rate步长0.25 rps；每台机器实验串行。'),
 ('评分','完整输出且 TTFT、TPOT 同时严格小于阈值才算达标。分母包含全部请求；2×边界两次各有1次有效120秒超时。'),
 ('边界规则','首个有效结果<90%即固定终点并确认一次；恰好90%继续。确认不与首测取平均或最佳值。基线没有继续搜索各自容量。'),
 ('主能耗','所有八张GPU的主窗口积分，包含100秒到达、实际排空及控制尾段。不是服务器墙上电表能耗。'),
 ('指标定义','goodput=达标请求数/实测主窗口秒数；J/good=八卡主能耗/达标请求数。TTFT、TPOT均值按完整请求计算。'),
 ('对比公式','节能率=1−PDBlend/基线；达标率差=PDBlend−基线，单位百分点；吞吐变化=PDBlend/基线−1。仅同机同rate首测配对。'),
 ('图表说明','每档6张可编辑XY数值曲线；主曲线只含首测。橙色菱形表示边界确认，90%目标线单独标示。'),
 ('解释范围','SLO倍率与机器绑定，跨倍率差异包含机器影响；单种子不提供独立种子置信区间。'),
 ('能耗分账','正式测量包含确认。准备能耗仅列已计量且不重叠窗口；未计量间隙保持未知，不将外层操作能耗重复累加。'),
 ('实现范围','仓库冻结机制复现；PDBlend为两个mixed TP1＋DVFS，PD拆分关闭。实例数、批次预算和原生实现的单独贡献未做消融。'),
]
for i,(label,text) in enumerate(notes,9):
    overview.write(i,0,label,fmt['section']);overview.merge_range(i,1,i,10,text,fmt['note']);overview.set_row(i,38)

headers=['系统','Rate (rps)','测量类型','重复编号','全部请求','完整输出请求','达标请求','SLO达标率','TTFT均值(s)','TPOT均值(ms)','Goodput(req/s)','八卡能耗(J)','八卡能耗(kJ)','J/达标请求','主窗口(s)','八卡平均功率(W)','TTFT要求(s)','TPOT要求(ms)','测量可计分','工作全部完成','达到90%目标','工程尝试','SLO倍率','机器','Trace SHA256','Cell ID']
locations={}
for host,label in [('A','0.5x'),('C','2x')]:
    ws=sheet(label+'测量明细',label+' SLO｜五系统逐次测量',f'机器 {host}；八卡主能耗包含到达和排空。黄色行为PDBlend边界确认，达标率低于90%显示红色。')
    rr=[r for r in rows if r['measurement_host']==host]
    data=[]
    for r in rr:
        data.append([names[r['system']],float(r['rate_rps']),'首测' if r['repeat']=='1' else '边界确认',int(r['repeat']),int(float(r['n_expected'])),int(float(r['completed_work_requests'])),int(float(r['good_requests'])),float(r['slo_attainment']),float(r['ttft_avg_s']),1000*float(r['tpot_avg_s']),float(r['goodput_measurement_rps']),float(r['energy_j']),float(r['energy_j'])/1000,float(r['energy_per_good_request_j']),float(r['measurement_duration_s']),float(r['energy_j'])/float(r['measurement_duration_s']),float(r['slo_ttft_s']),1000*float(r['slo_tpot_s']),'是' if r['measurement_valid']=='True' else '否','是' if r['work_complete']=='True' else '否','是' if r['slo_pass']=='True' else '否',int(r['engineering_attempt']),float(r['slo_scale']),host,r['trace_sha256'],r['cell_id']])
    table(ws,headers,data,{1:'two',7:'pct',8:'num',9:'num',10:'num',11:'num',12:'num',13:'two',14:'num',15:'two',16:'two',17:'two',22:'two'})
    for i,r in enumerate(rr,5):
        locations[r['cell_id']]=(ws.get_name(),i)
        for col,formula in [(7,f'=G{i+1}/E{i+1}'),(10,f'=G{i+1}/O{i+1}'),(12,f'=L{i+1}/1000'),(13,f'=L{i+1}/G{i+1}'),(15,f'=L{i+1}/O{i+1}')]:
            ws.write_formula(i,col,formula,fmt['pct' if col==7 else 'two' if col in (13,15) else 'num'],data[i-5][col])
        if r['repeat']=='2':ws.set_row(i,24,fmt['confirm'])
    ws.conditional_format(5,7,4+len(rr),7,{'type':'cell','criteria':'<','value':.9,'format':fmt['bad']})
    ws.conditional_format(5,7,4+len(rr),7,{'type':'cell','criteria':'>=','value':.9,'format':fmt['good']})
    ws.conditional_format(5,0,4+len(rr),25,{'type':'formula','criteria':'=$D6=2','format':fmt['confirm']})
    ws.set_column(24,25,30)

pair_count=0
for host,label in [('A','0.5x'),('C','2x')]:
    ws=sheet(label+'基线对比',label+' SLO｜PDBlend 相对四个基线','仅同机同rate首测配对，确认不混入。节能率为正表示PDBlend更省能；达标率差为正表示PDBlend更高。')
    rr=[r for r in rows if r['measurement_host']==host and r['repeat']=='1']
    ps=[r for r in rr if r['system']=='pdblend'];data=[];refs=[]
    for p in ps:
        for sysname in systems[1:]:
            b=next(r for r in rr if r['system']==sysname and r['rate_rps']==p['rate_rps'])
            pe,be=float(p['energy_j']),float(b['energy_j']);pg,bg=float(p['energy_per_good_request_j']),float(b['energy_per_good_request_j']);pa,ba=float(p['slo_attainment']),float(b['slo_attainment']);pt,bt=float(p['goodput_measurement_rps']),float(b['goodput_measurement_rps'])
            data.append([float(p['rate_rps']),names[sysname],pa,ba,(pa-ba)*100,pe/1000,be/1000,1-pe/be,pg,bg,1-pg/bg,pt,bt,pt/bt-1,'是' if pa>=.9 else '否',p['trace_sha256']])
            refs.append((locations[p['cell_id']],locations[b['cell_id']]))
    ph=['Rate (rps)','对照基线','PDB达标率','基线达标率','达标率差(百分点)','PDB能耗(kJ)','基线能耗(kJ)','总能耗降低','PDB J/good','基线 J/good','J/good降低','PDB Goodput','基线 Goodput','Goodput变化','PDB达到90%','Trace SHA256']
    table(ws,ph,data,{0:'two',2:'pct',3:'pct',4:'two',5:'num',6:'num',7:'pct',8:'two',9:'two',10:'pct',11:'num',12:'num',13:'pct'})
    for i,((sn,pr),(_,br)) in enumerate(refs,5):
        def ref(r,c):return f"'{sn}'!{cell(r,c)}"
        formulas={2:f'={ref(pr,7)}',3:f'={ref(br,7)}',4:f'=(C{i+1}-D{i+1})*100',5:f'={ref(pr,12)}',6:f'={ref(br,12)}',7:f'=1-F{i+1}/G{i+1}',8:f'={ref(pr,13)}',9:f'={ref(br,13)}',10:f'=1-I{i+1}/J{i+1}',11:f'={ref(pr,10)}',12:f'={ref(br,10)}',13:f'=L{i+1}/M{i+1}-1'}
        for c,f in formulas.items():ws.write_formula(i,c,f,fmt['pct' if c in [2,3,7,10,13] else 'two' if c in [4,8,9] else 'num'],data[i-5][c])
    ws.conditional_format(5,4,4+len(data),4,{'type':'3_color_scale','min_color':'F9CECE','mid_color':'FFFFFF','max_color':'BFE3CF','mid_type':'num','mid_value':0})
    pair_count+=len(data)

metrics=[('SLO达标率','slo_attainment',1,7,'0%'),('TTFT均值 (s)','ttft_avg_s',1,8,'0.0'),('TPOT均值 (ms)','tpot_avg_s',1000,9,'0'),('Goodput (req/s)','goodput_measurement_rps',1,10,'0.00'),('八卡能耗 (kJ)','energy_j',.001,12,'0'),('每个达标请求能耗 (J)','energy_per_good_request_j',1,13,'0')]
for host,label in [('A','0.5x'),('C','2x')]:
    ws=sheet(label+'曲线',label+' SLO｜六项指标曲线',f'机器 {host}；横轴为数值rate。实线与圆点=首测；橙色菱形=PDBlend边界确认。各项指标的图表数据位于右侧AD列起。')
    ws.freeze_panes(3,0);ws.set_column(0,21,9)
    rr=[r for r in rows if r['measurement_host']==host];rates=sorted({float(r['rate_rps']) for r in rr})
    lookup={(r['system'],float(r['rate_rps'])):r for r in rr if r['repeat']=='1'}
    confirm=next(r for r in rr if r['repeat']=='2')
    for mi,(title,key,scale,measurecol,nf) in enumerate(metrics):
        top=4+mi*13;left=29
        ws.write_row(top,left,['Rate']+[names[s] for s in systems]+['90%目标','确认Rate','确认值'],fmt['head'])
        for j,rate in enumerate(rates,top+1):
            ws.write_number(j,left,rate,fmt['two'])
            for k,s in enumerate(systems,1):
                r=lookup[s,rate];sn,rn=locations[r['cell_id']]
                ws.write_formula(j,left+k,f"='{sn}'!{cell(rn,measurecol)}",fmt['pct' if key=='slo_attainment' else 'num'],float(r[key])*scale)
            if key=='slo_attainment':ws.write_number(j,left+6,.9,fmt['pct'])
        ws.write_number(top+1,left+7,float(confirm['rate_rps']),fmt['two'])
        sn,rn=locations[confirm['cell_id']]
        ws.write_formula(top+1,left+8,f"='{sn}'!{cell(rn,measurecol)}",fmt['pct' if key=='slo_attainment' else 'num'],float(confirm[key])*scale)
        chart=wb.add_chart({'type':'scatter','subtype':'straight_with_markers'})
        for k,s in enumerate(systems,1):
            chart.add_series({'name':names[s],'categories':[ws.get_name(),top+1,left,top+len(rates),left],'values':[ws.get_name(),top+1,left+k,top+len(rates),left+k],
                'line':{'color':colors[s],'width':2.5 if s=='pdblend' else 1.5},'marker':{'type':'circle','size':5,'border':{'color':colors[s]},'fill':{'color':colors[s]}}})
        if key=='slo_attainment':
            chart.add_series({'name':'90%目标','categories':[ws.get_name(),top+1,left,top+len(rates),left],'values':[ws.get_name(),top+1,left+6,top+len(rates),left+6],'line':{'color':'888888','dash_type':'dash'},'marker':{'type':'none'}})
        chart.add_series({'name':'PDBlend确认','categories':[ws.get_name(),top+1,left+7,top+1,left+7],'values':[ws.get_name(),top+1,left+8,top+1,left+8],'line':{'none':True},'marker':{'type':'diamond','size':8,'border':{'color':'C96614'},'fill':{'color':'F5AE56'}}})
        chart.set_title({'name':title,'name_font':{'name':'Microsoft YaHei','size':12}})
        chart.set_x_axis({'name':'Rate (rps)','min':0,'max':rates[-1]+.125,'major_unit':.25,'num_format':'0.00','name_font':{'size':10}})
        ya={'num_format':nf,'min':0,'major_gridlines':{'visible':True,'line':{'color':'E6EBF0'}}}
        if key=='slo_attainment':ya.update({'max':1.05,'major_unit':.1})
        chart.set_y_axis(ya);chart.set_legend({'position':'bottom','font':{'name':'Microsoft YaHei','size':8}})
        chart.set_chartarea({'border':{'none':True}});chart.set_style(10);chart.set_size({'width':690,'height':355})
        ws.insert_chart(4+(mi//2)*19,(mi%2)*11,chart)
    ws.print_area(0,0,60,21);ws.fit_to_pages(1,2)

ws=sheet('PDBlend失分诊断','PDBlend｜达标与失分分类','覆盖两档全部PDBlend测量，包含边界确认。完整请求延迟超标与未完整输出的有效超时分列。')
dh=['机器','SLO倍率','Rate (rps)','重复编号','全部请求','完整输出','达标','仅TTFT超标','仅TPOT超标','两者均超标','有效超时','容量拒绝','达标率','TTFT P50(s)','TTFT P95(s)','TPOT P50(ms)','TPOT P95(ms)','Cell ID']
dd=[]
for d in diag:
    dd.append([d['node'],float(d['slo_scale']),float(d['rate_rps']),int(d['repeat']),int(d['N']),int(d['full_complete']),int(d['good']),int(d['complete_ttft_only_fail']),int(d['complete_tpot_only_fail']),int(d['complete_both_fail']),int(d['incomplete_timeout']),int(d['incomplete_capacity']),float(d['slo_attainment']),float(d['complete_ttft_p50_s']),float(d['complete_ttft_p95_s']),1000*float(d['complete_tpot_p50_s']),1000*float(d['complete_tpot_p95_s']),d['cell_id']])
table(ws,dh,dd,{1:'two',2:'two',12:'pct',13:'num',14:'num',15:'num',16:'num'})

ws=sheet('能耗分账','八卡能耗｜正式测量与准备分账','正式测量包含边界确认；已计量准备窗口与主测量互不重叠。存在未计量间隙，整个实验总能耗未知。')
ed=[]
for host in ['A','C']:
    n=energy['nodes'][host]
    for s in systems:
        se=n['by_system'][s];ed.append([host,'正式测量',names[s],se['cell_attempts'],se['primary_energy_j']/1000,'完整主窗口；PDBlend含确认'])
    ed.append([host,'正式测量小计','全部系统',n['verified_primary_windows'],n['primary_energy_subtotal_j']/1000,'小计行，不与上述分项再次相加'])
    ed.append([host,'已计量准备小计','部署/资格/失败准备',n['setup']['measured_operation_count'],n['setup']['measured_subtotal_j']/1000,'仅已知且不重叠窗口'])
    ed.append([host,'完整实验总能耗','全部阶段','',None,'未知；存在未计量间隙'])
table(ws,['机器','计量范围','系统或阶段','次数','八卡能耗(kJ)','说明'],ed,{4:'num'})
ws.set_column(5,5,52)

ws=sheet('证据索引','测量证据与 SHA256','每次测量的请求、功率、checkpoint、receipt、summary、binding及汇总来源。导出时已核对所有列出的SHA256；文件路径指向实验工作区。')
table(ws,['SLO倍率','机器','系统','Rate (rps)','重复编号','证据类型','原始文件路径','SHA256','核对状态','Cell ID'],evidence,{0:'two',3:'two'})
ws.set_column(6,6,65);ws.set_column(7,7,68);ws.set_column(9,9,35)

ws=sheet('原始观测','62 次原始观测｜全部CSV字段','保留源 observations.csv 的全部列，限定A/0.5×与C/2×。JSON引用与原始状态字符串完整保留；空字段保持空白。')
rawheaders=list(rows[0]);numeric={'engineering_attempt','slo_scale','rate_rps','repeat','slo_attainment','slo_ttft_s','slo_tpot_s','ttft_avg_s','tpot_avg_s','goodput_measurement_rps','energy_j','energy_per_good_request_j','gpu_util','energy_measured_gpu_count','n_expected','completed_work_requests','good_requests','completion_fraction','measurement_duration_s','measurement_start_s','measurement_end_s','full_operation_energy_j','seed','arrival_window_s'}
rawdata=[[float(r[k]) if k in numeric and r[k] else r[k] for k in rawheaders] for r in rows]
table(ws,rawheaders,rawdata);ws.set_column(0,len(rawheaders)-1,25)

overview.write(22,0,'工作表导航',fmt['section'])
for i,name in enumerate(list(worksheets)[1:],23):
    overview.write_url(i,0,f"internal:'{name}'!A1",fmt['link'],name)
overview.write(35,0,'原始报告',fmt['section']);overview.merge_range(35,1,36,10,str(FINAL/'REPORT.md'),fmt['note'])
overview.activate();overview.set_first_sheet()
wb.close()

# Independent reader checks stored numeric/formula caches and chart structure.
import openpyxl
from zipfile import ZipFile
book=openpyxl.load_workbook(DEST,data_only=True)
formula_book=openpyxl.load_workbook(DEST,data_only=False)
assert len(book.sheetnames)==11
for r in rows:
    sn,i=locations[r['cell_id']];ws=book[sn]
    for col,expected in [(7,float(r['slo_attainment'])),(10,float(r['goodput_measurement_rps'])),(12,float(r['energy_j'])/1000),(13,float(r['energy_per_good_request_j']))]:
        assert math.isclose(ws.cell(i+1,col+1).value,expected,rel_tol=1e-12),r['cell_id']
assert pair_count==48
for host,label in [('A','0.5x'),('C','2x')]:
    compare=book[label+'基线对比']
    for row in compare.iter_rows(min_row=6,values_only=True):
        if row[0] is None:continue
        assert math.isclose(row[7],1-row[5]/row[6],abs_tol=1e-12)
        assert math.isclose(row[10],1-row[8]/row[9],abs_tol=1e-12)
        assert math.isclose(row[4],100*(row[2]-row[3]),abs_tol=1e-12)
    assert len(formula_book[label+'曲线']._charts)==6
with ZipFile(DEST) as z:
    charts=[n for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')]
    assert len(charts)==12
    assert all(b'scatterChart' in z.read(c) for c in charts)
    assert all(b'#REF!' not in z.read(n) for n in z.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'))
result={'workbook':str(DEST),'sha256':hashlib.sha256(DEST.read_bytes()).hexdigest(),'sheets':book.sheetnames,'formal_measurements':62,'A_measurements':21,'C_measurements':41,'boundary_confirmations':2,'paired_baseline_comparisons':48,'editable_xy_charts':12,'evidence_rows':len(evidence),'verified_unique_sources':len(sources),'validation':'passed','source_sha256':sources}
(OUT/'validation-manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in result.items() if k!='source_sha256'},ensure_ascii=False,indent=2))
